"""Build the linked HR attendance workbook (نظام الحضور والانصراف المربوط).

Usage: python build_workbook.py SOURCE.xlsx OUTPUT.xlsx

Employees, shifts and settings values are carried over from SOURCE; every
other sheet is rebuilt so that the employee code is the single key that links
attendance, leaves, missions, payroll, the employee report and the dashboard.
"""
import sys
import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]

EMP_ROWS = 500      # employees rows 2..501
ATT_ROWS = 3000     # attendance rows 2..3001
LV_ROWS = 1000      # leaves / missions rows 2..1001
SH_ROWS = 20        # shifts rows 2..21

EMP_LAST = EMP_ROWS + 1
ATT_LAST = ATT_ROWS + 1
LV_LAST = LV_ROWS + 1
SH_LAST = SH_ROWS + 1

S_SET, S_SHIFT, S_EMP, S_ATT = "الإعدادات", "الورديات", "الموظفين", "الحضور والانصراف"
S_LV, S_MS, S_PAY, S_REP, S_DASH, S_HELP = (
    "الإجازات", "المأموريات", "تقرير الرواتب", "تقرير موظف", "Dashboard", "طريقة الاستخدام")


def q(name):
    return f"'{name}'"


EMP = f"{q(S_EMP)}!$A$2:$J${EMP_LAST}"
ATT = q(S_ATT)
SET = q(S_SET)

DAYS = ["الأحد", "الإثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت"]
DAYS_ARR = "{" + ",".join(f'"{d}"' for d in DAYS) + "}"
CHOOSE_DAY = 'CHOOSE(WEEKDAY({c}),' + ",".join(f'"{d}"' for d in DAYS) + ")"

# ---------------------------------------------------------------- styles
FONT = "Arial"
F_BASE = Font(name=FONT, size=11)
F_HEAD = Font(name=FONT, size=11, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color="1F3864")
F_BOLD = Font(name=FONT, size=11, bold=True)
F_INPUT = Font(name=FONT, size=11, color="0000FF")
F_LINK = Font(name=FONT, size=11, color="008000")
FILL_HEAD_IN = PatternFill("solid", fgColor="C55A11")    # input column header
FILL_HEAD_CALC = PatternFill("solid", fgColor="1F3864")  # formula column header
FILL_IN = PatternFill("solid", fgColor="FFF2CC")         # cell the user types in
FILL_KPI = PatternFill("solid", fgColor="DEEAF6")
FILL_TOTAL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

FMT_DATE = "yyyy-mm-dd"
FMT_TIME = "hh:mm"
FMT_MONEY = "#,##0.00"
FMT_NUM = "#,##0.##"
FMT_MONTH = "mmmm yyyy"


def new_sheet(wb, title, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def header(ws, row, cols):
    """cols: list of (title, is_input)."""
    for i, (title, is_input) in enumerate(cols, 1):
        c = ws.cell(row, i, title)
        c.font = F_HEAD
        c.fill = FILL_HEAD_IN if is_input else FILL_HEAD_CALC
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 32


def fill_table(ws, first, last, cols):
    """cols: list of (formula_template or None, number_format, is_input).
    formula_template uses {r} for the row number."""
    for r in range(first, last + 1):
        for i, (tpl, fmt, is_input) in enumerate(cols, 1):
            c = ws.cell(r, i)
            if tpl is not None:
                c.value = tpl.format(r=r)
            c.font = F_INPUT if is_input else F_BASE
            if is_input:
                c.fill = FILL_IN
            if fmt:
                c.number_format = fmt
            c.border = BORDER
            c.alignment = CENTER


def add_list(ws, ref, formula1, title=None, prompt=None):
    dv = DataValidation(type="list", formula1=formula1, allow_blank=True,
                        showErrorMessage=True, errorTitle="قيمة غير صحيحة",
                        error="اختر قيمة من القائمة")
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = title or "", prompt, True
    ws.add_data_validation(dv)
    dv.add(ref)


def label(ws, ref, text, bold=True):
    c = ws[ref]
    c.value = text
    c.font = F_BOLD if bold else F_BASE
    c.alignment = RIGHT
    c.border = BORDER
    c.fill = FILL_TOTAL


def to_time(v):
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, str) and ":" in v:
        h, m = v.strip().split(":")[:2]
        return dt.time(int(h) % 24, int(m))
    return v


# ---------------------------------------------------------------- source data
src = openpyxl.load_workbook(SRC)
src_emp = [
    [c.value for c in row]
    for row in src[S_EMP].iter_rows(min_row=2, max_col=10)
    if row[0].value not in (None, "")
]
src_shifts = [
    [c.value for c in row]
    for row in src[S_SHIFT].iter_rows(min_row=2, max_col=6)
    if row[0].value not in (None, "")
]
src_set = {r[0].value: r[1].value for r in src[S_SET].iter_rows(min_row=2, max_col=2)}

wb = openpyxl.Workbook()
wb.remove(wb.active)

# ================================================================ instructions
ws = new_sheet(wb, S_HELP, [4, 30, 90])
ws["B1"] = "نظام الحضور والانصراف والرواتب — دليل الاستخدام"
ws["B1"].font = F_TITLE
rows = [
    ("الفكرة", "كود الموظف هو المفتاح الوحيد الذي يربط كل الشيتات. سجّل الموظف مرة واحدة في شيت «الموظفين»، وبعدها أي تسجيل حضور أو إجازة أو مأمورية بالكود يظهر تلقائياً في كل التقارير."),
    ("1) الإعدادات", "حدّد شهر التقارير، ساعات العمل اليومية، أيام الشهر لحساب الراتب، يوم الراحة الأسبوعية والحد الأدنى للإضافي."),
    ("2) الورديات", "اسم كل وردية ومواعيد بدايتها ونهايتها ودقائق السماح بالتأخير. الورديات الليلية التي تعبر منتصف الليل محسوبة صح."),
    ("3) الموظفين", "اكتب: الكود، الاسم، القسم، الوظيفة، اختر الوردية، قيمة اليوم، قيمة ساعة الإضافي، والحالة. مواعيد الدوام تُسحب تلقائياً من الوردية."),
    ("4) الحضور والانصراف", "لكل يوم: اكتب التاريخ واختر كود الموظف واكتب وقت الحضور والانصراف (مثال 08:05 و 16:30). الاسم والقسم والوردية وساعات العمل والتأخير والإضافي والحالة والخصومات تُحسب تلقائياً."),
    ("5) الإجازات", "سجّل الإجازة بالكود ومن/إلى تاريخ واجعل الحالة «مقبولة» — تظهر تلقائياً في عمود الإجازة بشيت الحضور وفي أيام الإجازة بتقرير الرواتب."),
    ("6) المأموريات", "مأموريات بالساعات داخل اليوم؛ الموافَق عليها تُجمع في تقرير الرواتب."),
    ("7) تقرير الرواتب", "يتولد تلقائياً لكل الموظفين عن الشهر المحدد في الإعدادات: الحضور، الغياب، الإجازات، الساعات، التأخير، الإضافي، الخصومات وصافي الراتب."),
    ("8) تقرير موظف", "اختر كود الموظف والشهر فتظهر بياناته وحركته يوماً بيوم وملخص الشهر."),
    ("9) Dashboard", "مؤشرات يوم محدد (افتراضياً اليوم) ومؤشرات الشهر."),
    ("دليل الألوان", "الخلايا الصفراء بخط أزرق = خانات إدخال تكتب فيها. رؤوس الأعمدة البرتقالية = إدخال، الكحلية = معادلات لا تُعدّل. الخط الأخضر = قيمة مسحوبة من شيت آخر."),
    ("مثال صف حضور", "التاريخ 2026-10-01 | الكود EMP001 | الحضور 08:20 | الانصراف 17:30  ←  التأخير 20 دقيقة (أكبر من السماح 15)، الإضافي 90 دقيقة، الحالة «متأخر»."),
    ("ملاحظات", "• حالة يدوية: اختر «غياب بعذر» لعدم خصم اليوم، أو «راحة أسبوعية»، أو «إجازة/مأمورية». • يوم الراحة الأسبوعية بدون حضور لا يُحسب غياباً. • «إجازة بدون راتب» و«غائب» يُخصم عنها قيمة اليوم. • خصم التأخير = (دقائق التأخير + الانصراف المبكر) ÷ 60 × أجر الساعة × معامل الخصم."),
]
for i, (k, v) in enumerate(rows, 3):
    ws.cell(i, 2, k).font = F_BOLD
    ws.cell(i, 3, v).font = F_BASE
    ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")
    ws.cell(i, 2).alignment = Alignment(vertical="top", horizontal="right")
    ws.row_dimensions[i].height = 48

# ================================================================ settings
ws = new_sheet(wb, S_SET, [36, 22, 60])
ws["A1"] = "الإعدادات العامة"
ws["A1"].font = F_TITLE
header(ws, 2, [("الإعداد", False), ("القيمة", True), ("شرح", False)])
settings = [
    # row, label, value, fmt, note, is_input
    (3, "شهر التقارير", dt.datetime(2026, 10, 1), FMT_MONTH, "اكتب أي تاريخ داخل الشهر المطلوب — يُستخدم في تقرير الرواتب والـ Dashboard", True),
    (4, "بداية الشهر", "=DATE(YEAR(B3),MONTH(B3),1)", FMT_DATE, "تلقائي", False),
    (5, "نهاية الشهر", "=DATE(YEAR(B3),MONTH(B3)+1,0)", FMT_DATE, "تلقائي", False),
    (6, "ساعات العمل اليومية", src_set.get("ساعات العمل اليومية", 8), FMT_NUM, "لحساب أجر الساعة = قيمة اليوم ÷ ساعات العمل (من ملفك الأصلي)", True),
    (7, "أيام الشهر لحساب الراتب", 30, FMT_NUM, "الراتب الأساسي = قيمة اليوم × هذا الرقم (افتراض: 30 يوم)", True),
    (8, "الحد الأدنى للإضافي (دقيقة)", (src_set.get("الحد الأدنى لساعات الإضافي", 1) or 0) * 60, FMT_NUM, "الإضافي الأقل من ذلك لا يُحتسب (من ملفك الأصلي: 1 ساعة)", True),
    (9, "معامل خصم التأخير", 1, FMT_NUM, "1 = خصم دقيقة بدقيقة، 2 = ضعف، 0 = بدون خصم تأخير (افتراض: 1)", True),
    (10, "يوم الراحة الأسبوعية", "الجمعة", None, "اختر من القائمة (افتراض: الجمعة)", True),
    (11, "رقم يوم الراحة", f"=MATCH(B10,{DAYS_ARR},0)", FMT_NUM, "تلقائي (الأحد = 1)", False),
    (12, "قيمة اليوم الافتراضية", src_set.get("قيمة يوم العمل", 100), FMT_MONEY, "مرجع فقط عند إضافة موظف جديد (من ملفك الأصلي)", True),
    (13, "قيمة ساعة الإضافي الافتراضية", src_set.get("قيمة ساعة الإضافي", 25), FMT_MONEY, "مرجع فقط عند إضافة موظف جديد (من ملفك الأصلي)", True),
]
for r, lab, val, fmt, note, is_in in settings:
    ws.cell(r, 1, lab).font = F_BOLD
    c = ws.cell(r, 2, val)
    c.font = F_INPUT if is_in else F_BASE
    if is_in:
        c.fill = FILL_IN
    if fmt:
        c.number_format = fmt
    ws.cell(r, 3, note).font = F_BASE
    for col in range(1, 4):
        ws.cell(r, col).border = BORDER
        ws.cell(r, col).alignment = RIGHT
add_list(ws, "B10", '"' + ",".join(DAYS) + '"')

SET_START, SET_END = f"{SET}!$B$4", f"{SET}!$B$5"
SET_HOURS, SET_MDAYS, SET_OTMIN = f"{SET}!$B$6", f"{SET}!$B$7", f"{SET}!$B$8"
SET_LATEX, SET_REST = f"{SET}!$B$9", f"{SET}!$B$11"

# ================================================================ shifts
ws = new_sheet(wb, S_SHIFT, [20, 16, 16, 14, 18, 30])
header(ws, 1, [("اسم الوردية", True), ("بداية الدوام", True), ("نهاية الدوام", True),
               ("ساعات العمل", False), ("سماح التأخير (د)", True), ("ملاحظات", True)])
fill_table(ws, 2, SH_LAST, [
    (None, None, True), (None, FMT_TIME, True), (None, FMT_TIME, True),
    ('=IF(OR(B{r}="",C{r}=""),"",ROUND(MOD(C{r}-B{r},1)*24,2))', FMT_NUM, False),
    (None, FMT_NUM, True), (None, None, True)])
for i, s in enumerate(src_shifts, 2):
    ws.cell(i, 1, s[0])
    ws.cell(i, 2, to_time(s[1]))
    ws.cell(i, 3, to_time(s[2]))
    ws.cell(i, 5, s[4] if s[4] is not None else 15)
    ws.cell(i, 6, s[5])
ws.freeze_panes = "A2"
SHIFTS = f"{q(S_SHIFT)}!$A$2:$E${SH_LAST}"

# ================================================================ employees
ws = new_sheet(wb, S_EMP, [14, 26, 18, 18, 14, 13, 13, 13, 16, 12, 16])
header(ws, 1, [("كود الموظف", True), ("اسم الموظف", True), ("القسم", True), ("الوظيفة", True),
               ("الوردية", True), ("بداية الدوام", False), ("نهاية الدوام", False),
               ("قيمة اليوم", True), ("قيمة ساعة الإضافي", True), ("الحالة", True),
               ("الراتب الأساسي", False)])
fill_table(ws, 2, EMP_LAST, [
    (None, None, True), (None, None, True), (None, None, True), (None, None, True),
    (None, None, True),
    (f'=IF(E{{r}}="","",IFERROR(VLOOKUP(E{{r}},{SHIFTS},2,FALSE),""))', FMT_TIME, False),
    (f'=IF(E{{r}}="","",IFERROR(VLOOKUP(E{{r}},{SHIFTS},3,FALSE),""))', FMT_TIME, False),
    (None, FMT_MONEY, True), (None, FMT_MONEY, True), (None, None, True),
    (f'=IF(OR(A{{r}}="",H{{r}}=""),"",H{{r}}*{SET_MDAYS})', FMT_MONEY, False)])
for i, e in enumerate(src_emp, 2):
    for col in (1, 2, 3, 4, 5, 8, 9, 10):
        ws.cell(i, col, e[col - 1])
for r in range(2, EMP_LAST + 1):
    for col in (6, 7):
        ws.cell(r, col).font = F_LINK
add_list(ws, f"E2:E{EMP_LAST}", f"={q(S_SHIFT)}!$A$2:$A${SH_LAST}")
add_list(ws, f"J2:J{EMP_LAST}", '"نشط,موقوف,مستقيل"')
ws.conditional_formatting.add(
    f"A2:A{EMP_LAST}",
    FormulaRule(formula=[f'AND(A2<>"",COUNTIF($A$2:$A${EMP_LAST},A2)>1)'],
                fill=PatternFill("solid", fgColor="FF9999")))
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:K{EMP_LAST}"

# ================================================================ leaves
ws = new_sheet(wb, S_LV, [14, 24, 20, 14, 14, 11, 16, 30, 18])
header(ws, 1, [("كود الموظف", True), ("اسم الموظف", False), ("نوع الإجازة", True),
               ("من تاريخ", True), ("إلى تاريخ", True), ("عدد الأيام", False),
               ("الحالة", True), ("ملاحظات", True), ("أيام داخل شهر التقارير", False)])
fill_table(ws, 2, LV_LAST, [
    (None, None, True),
    (f'=IF(A{{r}}="","",IFERROR(VLOOKUP(A{{r}},{EMP},2,FALSE),"⚠ كود غير موجود"))', None, False),
    (None, None, True), (None, FMT_DATE, True), (None, FMT_DATE, True),
    ('=IF(OR(D{r}="",E{r}=""),"",E{r}-D{r}+1)', "0", False),
    (None, None, True), (None, None, True),
    (f'=IF(OR(D{{r}}="",E{{r}}=""),"",MAX(0,MIN(E{{r}},{SET_END})-MAX(D{{r}},{SET_START})+1))', "0", False)])
for r in range(2, LV_LAST + 1):
    ws.cell(r, 2).font = F_LINK
LEAVE_TYPES = "إجازة سنوية,إجازة مرضية,إجازة اضطرارية,إجازة بدون راتب,مأمورية,مهمة رسمية,أخرى"
add_list(ws, f"A2:A{LV_LAST}", f"={q(S_EMP)}!$A$2:$A${EMP_LAST}")
add_list(ws, f"C2:C{LV_LAST}", f'"{LEAVE_TYPES}"')
add_list(ws, f"G2:G{LV_LAST}", '"مقبولة,قيد المراجعة,مرفوضة"')
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:I{LV_LAST}"
LV = q(S_LV)

# ================================================================ missions
ws = new_sheet(wb, S_MS, [14, 24, 14, 12, 12, 11, 16, 16, 30])
header(ws, 1, [("كود الموظف", True), ("اسم الموظف", False), ("التاريخ", True),
               ("من الساعة", True), ("إلى الساعة", True), ("الساعات", False),
               ("نوع المأمورية", True), ("الموافقة", True), ("ملاحظات", True)])
fill_table(ws, 2, LV_LAST, [
    (None, None, True),
    (f'=IF(A{{r}}="","",IFERROR(VLOOKUP(A{{r}},{EMP},2,FALSE),"⚠ كود غير موجود"))', None, False),
    (None, FMT_DATE, True), (None, FMT_TIME, True), (None, FMT_TIME, True),
    ('=IF(OR(D{r}="",E{r}=""),"",ROUND(MOD(E{r}-D{r},1)*24,2))', FMT_NUM, False),
    (None, None, True), (None, None, True), (None, None, True)])
for r in range(2, LV_LAST + 1):
    ws.cell(r, 2).font = F_LINK
add_list(ws, f"A2:A{LV_LAST}", f"={q(S_EMP)}!$A$2:$A${EMP_LAST}")
add_list(ws, f"G2:G{LV_LAST}", '"ميدانية,إدارية,موقع,أخرى"')
add_list(ws, f"H2:H{LV_LAST}", '"موافق,قيد المراجعة,مرفوض"')
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:I{LV_LAST}"
MS = q(S_MS)

# ================================================================ attendance
ws = new_sheet(wb, S_ATT, [13, 13, 24, 16, 12, 11, 11, 11, 11, 11, 11, 11, 11,
                           18, 16, 18, 12, 14, 12, 26, 4])
header(ws, 1, [("التاريخ", True), ("كود الموظف", True), ("اسم الموظف", False), ("القسم", False),
               ("الوردية", False), ("بداية الدوام", False), ("نهاية الدوام", False),
               ("الحضور", True), ("الانصراف", True), ("ساعات العمل", False),
               ("التأخير (د)", False), ("انصراف مبكر (د)", False), ("إضافي (د)", False),
               ("إجازة/مأمورية (تلقائي)", False), ("حالة يدوية", True), ("الحالة", False),
               ("خصم الغياب", False), ("خصم التأخير", False), ("قيمة الإضافي", False),
               ("ملاحظات", True), ("مفتاح", False)])

emp_lookup = 'IF($B{r}="","",IFERROR(VLOOKUP($B{r},' + EMP + ',{n},FALSE){txt},{miss}))'
late_raw = "ROUND(MAX(0,MOD(H{r}-F{r}+0.5,1)-0.5)*1440,0)"
early_raw = "ROUND(MAX(0,MOD(G{r}-I{r}+0.5,1)-0.5)*1440,0)"
ot_raw = "ROUND(MAX(0,MOD(I{r}-G{r}+0.5,1)-0.5)*1440,0)"
grace = f"IFERROR(VLOOKUP(E{{r}},{SHIFTS},5,FALSE),0)"
leave_match = (f"({LV}!$A$2:$A${LV_LAST}=B{{r}})*({LV}!$D$2:$D${LV_LAST}<=A{{r}})"
               f"*({LV}!$E$2:$E${LV_LAST}>=A{{r}})*({LV}!$G$2:$G${LV_LAST}=\"مقبولة\")")
leave_count = (f'COUNTIFS({LV}!$A$2:$A${LV_LAST},B{{r}},{LV}!$D$2:$D${LV_LAST},"<="&A{{r}},'
               f'{LV}!$E$2:$E${LV_LAST},">="&A{{r}},{LV}!$G$2:$G${LV_LAST},"مقبولة")')
daily = f"IFERROR(VLOOKUP(B{{r}},{EMP},8,FALSE),0)"
ot_rate = f"IFERROR(VLOOKUP(B{{r}},{EMP},9,FALSE),0)"

att_cols = [
    (None, FMT_DATE, True),
    (None, None, True),
    ("=" + emp_lookup.format(r="{r}", n=2, miss='"⚠ كود غير موجود"', txt='&""'), None, False),
    ("=" + emp_lookup.format(r="{r}", n=3, miss='""', txt='&""'), None, False),
    ("=" + emp_lookup.format(r="{r}", n=5, miss='""', txt='&""'), None, False),
    ("=" + emp_lookup.format(r="{r}", n=6, miss='""', txt=""), FMT_TIME, False),
    ("=" + emp_lookup.format(r="{r}", n=7, miss='""', txt=""), FMT_TIME, False),
    (None, FMT_TIME, True),
    (None, FMT_TIME, True),
    ('=IF(OR(H{r}="",I{r}=""),"",ROUND(MOD(I{r}-H{r},1)*24,2))', FMT_NUM, False),
    (f'=IF(OR(H{{r}}="",F{{r}}=""),"",IF({late_raw}>{grace},{late_raw},0))', "0", False),
    (f'=IF(OR(I{{r}}="",G{{r}}=""),"",{early_raw})', "0", False),
    (f'=IF(OR(I{{r}}="",G{{r}}=""),"",IF({ot_raw}>={SET_OTMIN},{ot_raw},0))', "0", False),
    (f'=IF(OR(A{{r}}="",B{{r}}=""),"",IF({leave_count}=0,"",'
     f'LOOKUP(2,1/({leave_match}),{LV}!$C$2:$C${LV_LAST})))', None, False),
    (None, None, True),
    (f'=IF(B{{r}}="","",IF(ISNA(MATCH(B{{r}},{q(S_EMP)}!$A$2:$A${EMP_LAST},0)),"⚠ كود غير موجود",'
     'IF(O{r}<>"",O{r},IF(N{r}<>"",N{r},IF(H{r}="",'
     f'IF(AND(A{{r}}<>"",WEEKDAY(A{{r}})={SET_REST}),"راحة أسبوعية","غائب"),'
     'IF(I{r}="","بدون انصراف",IF(N(K{r})>0,"متأخر","حاضر")))))))', None, False),
    (f'=IF(B{{r}}="","",IF(OR(P{{r}}="غائب",P{{r}}="غياب",P{{r}}="إجازة بدون راتب"),{daily},0))',
     FMT_MONEY, False),
    (f'=IF(B{{r}}="","",ROUND((N(K{{r}})+N(L{{r}}))/60*{daily}/{SET_HOURS}*{SET_LATEX},2))',
     FMT_MONEY, False),
    (f'=IF(OR(B{{r}}="",M{{r}}=""),"",ROUND(M{{r}}/60*{ot_rate},2))', FMT_MONEY, False),
    (None, None, True),
    ('=IF(OR(A{r}="",B{r}=""),"",A{r}&"|"&B{r})', None, False),
]
fill_table(ws, 2, ATT_LAST, att_cols)
for r in range(2, ATT_LAST + 1):
    for col in (3, 4, 5, 6, 7, 14):
        ws.cell(r, col).font = F_LINK
add_list(ws, f"B2:B{ATT_LAST}", f"={q(S_EMP)}!$A$2:$A${EMP_LAST}",
         "كود الموظف", "اختر الكود — الاسم والقسم والوردية تظهر تلقائياً")
add_list(ws, f"O2:O{ATT_LAST}", '"إجازة,مأمورية,راحة أسبوعية,غياب بعذر,غياب"')
dv_time = DataValidation(type="time", operator="between", formula1="0", formula2="0.999988426",
                         allow_blank=True, showErrorMessage=True, errorTitle="وقت غير صحيح",
                         error="اكتب الوقت بصيغة 08:30")
ws.add_data_validation(dv_time)
dv_time.add(f"H2:I{ATT_LAST}")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         showErrorMessage=True, errorTitle="تاريخ غير صحيح",
                         error="اكتب التاريخ بصيغة 2026-10-01")
ws.add_data_validation(dv_date)
dv_date.add(f"A2:A{ATT_LAST}")

status_rng = f"P2:P{ATT_LAST}"
for txt, color in [("حاضر", "C6EFCE"), ("متأخر", "FFEB9C"), ("غائب", "FFC7CE"),
                   ("بدون انصراف", "F8CBAD"), ("راحة أسبوعية", "E7E6E6")]:
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                         fill=PatternFill("solid", fgColor=color)))
ws.conditional_formatting.add(status_rng, FormulaRule(
    formula=['AND(P2<>"",OR(LEFT(P2,5)="إجازة",LEFT(P2,6)="مأمورية",P2="مهمة رسمية",P2="غياب بعذر"))'],
    fill=PatternFill("solid", fgColor="BDD7EE")))
ws.conditional_formatting.add(f"C2:C{ATT_LAST}", FormulaRule(
    formula=['LEFT(C2,1)="⚠"'], font=Font(name=FONT, color="C00000", bold=True)))
ws.conditional_formatting.add(f"A2:T{ATT_LAST}", FormulaRule(
    formula=[f'AND($U2<>"",COUNTIF($U$2:$U${ATT_LAST},$U2)>1)'],
    fill=PatternFill("solid", fgColor="FF9999")))
ws.column_dimensions["U"].hidden = True
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:T{ATT_LAST}"


def att(col):
    return f"{ATT}!${col}$2:${col}${ATT_LAST}"


# ================================================================ payroll report
ws = new_sheet(wb, S_PAY, [13, 24, 16, 11, 11, 11, 11, 12, 12, 12, 12, 12, 14, 13, 13, 13, 15])
ws.merge_cells("A1:Q1")
ws["A1"] = "تقرير الحضور والرواتب الشهري"
ws["A1"].font = F_TITLE
ws["A1"].alignment = CENTER
label(ws, "A3", "الشهر")
ws["B3"] = f"={SET}!B3"
ws["B3"].number_format = FMT_MONTH
ws["B3"].font = F_LINK
label(ws, "C3", "من")
ws["D3"] = f"={SET}!B4"
label(ws, "E3", "إلى")
ws["F3"] = f"={SET}!B5"
for ref in ("D3", "F3"):
    ws[ref].number_format = FMT_DATE
    ws[ref].font = F_LINK
ws["H3"] = "لتغيير الشهر عدّل «شهر التقارير» في شيت الإعدادات"
ws["H3"].font = Font(name=FONT, italic=True, color="7F7F7F")
header(ws, 5, [("كود الموظف", False), ("اسم الموظف", False), ("القسم", False),
               ("أيام الحضور", False), ("أيام التأخير", False), ("أيام الغياب", False),
               ("أيام الإجازة", False), ("ساعات العمل", False), ("التأخير (د)", False),
               ("انصراف مبكر (د)", False), ("ساعات الإضافي", False), ("ساعات المأموريات", False),
               ("الراتب الأساسي", False), ("قيمة الإضافي", False), ("خصم الغياب", False),
               ("خصم التأخير", False), ("صافي الراتب", False)])
inm = f'{att("A")},">="&$D$3,{att("A")},"<="&$F$3'
pay_first, pay_last = 6, 6 + EMP_ROWS - 1


def by_emp(col, extra=""):
    return f'SUMIFS({att(col)},{att("B")},$A{{r}},{inm}{extra})'


def cnt(status):
    return f'COUNTIFS({att("B")},$A{{r}},{att("P")},"{status}",{inm})'


g = '=IF($A{r}="","",'
pay_cols = [
    (f'=IF({q(S_EMP)}!A{{e}}="","",{q(S_EMP)}!A{{e}})', None),
    (g + f'{q(S_EMP)}!B{{e}})', None),
    (g + f'IF({q(S_EMP)}!C{{e}}="","",{q(S_EMP)}!C{{e}}))', None),
    (g + cnt("حاضر") + "+" + cnt("متأخر") + "+" + cnt("بدون انصراف") + ")", "0"),
    (g + cnt("متأخر") + ")", "0"),
    (g + cnt("غائب") + "+" + cnt("غياب") + ")", "0"),
    (g + f'SUMIFS({LV}!$I$2:$I${LV_LAST},{LV}!$A$2:$A${LV_LAST},$A{{r}},{LV}!$G$2:$G${LV_LAST},"مقبولة"))', "0"),
    (g + by_emp("J") + ")", FMT_NUM),
    (g + by_emp("K") + ")", "0"),
    (g + by_emp("L") + ")", "0"),
    (g + f'ROUND({by_emp("M")}/60,2))', FMT_NUM),
    (g + f'SUMIFS({MS}!$F$2:$F${LV_LAST},{MS}!$A$2:$A${LV_LAST},$A{{r}},{MS}!$C$2:$C${LV_LAST},">="&$D$3,'
     f'{MS}!$C$2:$C${LV_LAST},"<="&$F$3,{MS}!$H$2:$H${LV_LAST},"موافق"))', FMT_NUM),
    (g + f'N({q(S_EMP)}!K{{e}}))', FMT_MONEY),
    (g + by_emp("S") + ")", FMT_MONEY),
    (g + by_emp("Q") + ")", FMT_MONEY),
    (g + by_emp("R") + ")", FMT_MONEY),
    (g + 'M{r}+N{r}-O{r}-P{r})', FMT_MONEY),
]
for i, r in enumerate(range(pay_first, pay_last + 1)):
    e = 2 + i
    for col, (tpl, fmt) in enumerate(pay_cols, 1):
        c = ws.cell(r, col, tpl.replace("{e}", str(e)).format(r=r))
        c.font = F_LINK if col <= 3 else F_BASE
        c.border = BORDER
        c.alignment = CENTER
        if fmt:
            c.number_format = fmt
tot = pay_last + 1
label(ws, f"A{tot}", "الإجمالي")
ws.merge_cells(f"A{tot}:C{tot}")
for col in range(4, 18):
    L = get_column_letter(col)
    c = ws.cell(tot, col, f"=SUM({L}{pay_first}:{L}{pay_last})")
    c.font = F_BOLD
    c.fill = FILL_TOTAL
    c.border = BORDER
    c.alignment = CENTER
    c.number_format = pay_cols[col - 1][1] or FMT_NUM
ws.freeze_panes = "C6"
ws.auto_filter.ref = f"A5:Q{pay_last}"
PAY_TOTAL_ROW = tot

# ================================================================ employee report
ws = new_sheet(wb, S_REP, [14, 12, 12, 12, 12, 12, 12, 18, 26])
ws.merge_cells("A1:I1")
ws["A1"] = "تقرير حضور موظف"
ws["A1"].font = F_TITLE
ws["A1"].alignment = CENTER
label(ws, "A3", "كود الموظف")
ws["B3"] = src_emp[0][0] if src_emp else ""
ws["B3"].font, ws["B3"].fill, ws["B3"].border = F_INPUT, FILL_IN, BORDER
add_list(ws, "B3", f"={q(S_EMP)}!$A$2:$A${EMP_LAST}", "كود الموظف", "اختر الكود")
label(ws, "C3", "الاسم")
ws.merge_cells("D3:F3")
ws["D3"] = f'=IF($B$3="","",IFERROR(VLOOKUP($B$3,{EMP},2,FALSE),"⚠ كود غير موجود"))'
label(ws, "G3", "الشهر")
ws.merge_cells("H3:I3")
ws["H3"] = f"={SET}!B3"
ws["H3"].number_format = FMT_MONTH
ws["H3"].font, ws["H3"].fill, ws["H3"].border = F_INPUT, FILL_IN, BORDER
label(ws, "A4", "القسم")
ws["B4"] = f'=IF($B$3="","",IFERROR(VLOOKUP($B$3,{EMP},3,FALSE)&"",""))'
label(ws, "C4", "الوردية")
ws["D4"] = f'=IF($B$3="","",IFERROR(VLOOKUP($B$3,{EMP},5,FALSE)&"",""))'
label(ws, "E4", "الدوام")
ws.merge_cells("F4:G4")
ws["F4"] = (f'=IF($B$3="","",IFERROR(TEXT(VLOOKUP($B$3,{EMP},6,FALSE),"hh:mm")&" - "&'
            f'TEXT(VLOOKUP($B$3,{EMP},7,FALSE),"hh:mm"),""))')
label(ws, "H4", "قيمة اليوم")
ws["I4"] = f'=IF($B$3="","",IFERROR(VLOOKUP($B$3,{EMP},8,FALSE),""))'
ws["I4"].number_format = FMT_MONEY
for ref in ("D3", "B4", "D4", "F4", "I4"):
    ws[ref].font, ws[ref].border, ws[ref].alignment = F_LINK, BORDER, CENTER
ws["H3"].alignment = CENTER
# hidden helpers: month start / end
ws["K3"] = "=DATE(YEAR($H$3),MONTH($H$3),1)"
ws["K4"] = "=DATE(YEAR($H$3),MONTH($H$3)+1,0)"
ws.column_dimensions["K"].hidden = True

# summary block rows 6-7
summary = [
    ("أيام الحضور", f'COUNTIF($H$11:$H$41,"حاضر")+COUNTIF($H$11:$H$41,"متأخر")+COUNTIF($H$11:$H$41,"بدون انصراف")', "0"),
    ("أيام الغياب", 'COUNTIF($H$11:$H$41,"غائب")+COUNTIF($H$11:$H$41,"غياب")', "0"),
    ("أيام الإجازة", 'COUNTIF($H$11:$H$41,"إجازة*")+COUNTIF($H$11:$H$41,"مأمورية*")+COUNTIF($H$11:$H$41,"مهمة رسمية")', "0"),
    ("ساعات العمل", "SUM($E$11:$E$41)", FMT_NUM),
    ("التأخير (د)", "SUM($F$11:$F$41)", "0"),
    ("ساعات الإضافي", "ROUND(SUM($G$11:$G$41)/60,2)", FMT_NUM),
    ("خصم الغياب", f'SUMIFS({att("Q")},{att("B")},$B$3,{att("A")},">="&$K$3,{att("A")},"<="&$K$4)', FMT_MONEY),
    ("خصم التأخير", f'SUMIFS({att("R")},{att("B")},$B$3,{att("A")},">="&$K$3,{att("A")},"<="&$K$4)', FMT_MONEY),
    ("قيمة الإضافي", f'SUMIFS({att("S")},{att("B")},$B$3,{att("A")},">="&$K$3,{att("A")},"<="&$K$4)', FMT_MONEY),
]
for i, (lab, f, fmt) in enumerate(summary):
    col = 1 + i
    h = ws.cell(6, col, lab)
    h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD_CALC, CENTER, BORDER
    c = ws.cell(7, col, f'=IF($B$3="","",{f})')
    c.font, c.fill, c.alignment, c.border, c.number_format = F_BOLD, FILL_KPI, CENTER, BORDER, fmt
ws.row_dimensions[6].height = 32
ws["A8"] = "صافي الراتب للشهر"
ws["A8"].font = F_BOLD
ws.merge_cells("A8:B8")
ws["C8"] = (f'=IF($B$3="","",N(IFERROR(VLOOKUP($B$3,{EMP},8,FALSE),0))*{SET_MDAYS}+N(I7)-N(G7)-N(H7))')
ws["C8"].number_format = FMT_MONEY
ws["C8"].font, ws["C8"].fill, ws["C8"].border = F_BOLD, FILL_KPI, BORDER
ws["D8"] = "(الراتب الأساسي + الإضافي − خصم الغياب − خصم التأخير)"
ws["D8"].font = Font(name=FONT, italic=True, color="7F7F7F")

header(ws, 10, [("التاريخ", False), ("اليوم", False), ("الحضور", False), ("الانصراف", False),
                ("ساعات العمل", False), ("التأخير (د)", False), ("إضافي (د)", False),
                ("الحالة", False), ("ملاحظات", False)])
key = '$A{r}&"|"&$B$3'
match = f"MATCH({key},{att('U')},0)"


def pick(col, blank='""'):
    return f'=IF(OR($A{{r}}="",$B$3=""),"",IFERROR(INDEX({att(col)},{match})&"",{blank}))'


def pick_num(col):
    return (f'=IF(OR($A{{r}}="",$B$3=""),"",IFERROR(IF(INDEX({att(col)},{match})="","",'
            f'INDEX({att(col)},{match})),""))')


rep_leave_count = leave_count.replace("B{r}", "$B$3").replace("A{r}", "$A{r}")
rep_leave_match = leave_match.replace("B{r}", "$B$3").replace("A{r}", "$A{r}")
rep_cols = [
    ('=IF({d}>DAY($K$4),"",DATE(YEAR($H$3),MONTH($H$3),{d}))', FMT_DATE),
    ('=IF($A{r}="","",' + CHOOSE_DAY.format(c="$A{r}") + ")", None),
    (pick_num("H"), FMT_TIME),
    (pick_num("I"), FMT_TIME),
    (pick_num("J"), FMT_NUM),
    (pick_num("K"), "0"),
    (pick_num("M"), "0"),
    (f'=IF(OR($A{{r}}="",$B$3=""),"",IFERROR(INDEX({att("P")},{match}),'
     f'IF({rep_leave_count}>0,LOOKUP(2,1/({rep_leave_match}),{LV}!$C$2:$C${LV_LAST}),'
     f'IF(WEEKDAY($A{{r}})={SET_REST},"راحة أسبوعية","لا يوجد تسجيل"))))', None),
    (pick("T"), None),
]
for d in range(1, 32):
    r = 10 + d
    for col, (tpl, fmt) in enumerate(rep_cols, 1):
        c = ws.cell(r, col, tpl.replace("{d}", str(d)).format(r=r))
        c.font, c.border, c.alignment = F_BASE, BORDER, CENTER
        if fmt:
            c.number_format = fmt
for txt, color in [("حاضر", "C6EFCE"), ("متأخر", "FFEB9C"), ("غائب", "FFC7CE"),
                   ("بدون انصراف", "F8CBAD"), ("راحة أسبوعية", "E7E6E6"), ("لا يوجد تسجيل", "F2F2F2")]:
    ws.conditional_formatting.add("H11:H41", CellIsRule(operator="equal", formula=[f'"{txt}"'],
                                                        fill=PatternFill("solid", fgColor=color)))
ws.conditional_formatting.add("H11:H41", FormulaRule(
    formula=['AND(H11<>"",OR(LEFT(H11,5)="إجازة",LEFT(H11,6)="مأمورية",H11="مهمة رسمية",H11="غياب بعذر"))'],
    fill=PatternFill("solid", fgColor="BDD7EE")))

# ================================================================ dashboard
ws = new_sheet(wb, S_DASH, [24, 16, 4, 24, 16, 4, 24, 16])
ws.merge_cells("A1:H1")
ws["A1"] = "لوحة تحكم الموارد البشرية — الحضور والانصراف"
ws["A1"].font = F_TITLE
ws["A1"].alignment = CENTER
label(ws, "A3", "يوم المتابعة")
ws["B3"] = "=TODAY()"
ws["B3"].number_format = FMT_DATE
ws["B3"].font, ws["B3"].fill, ws["B3"].border = F_INPUT, FILL_IN, BORDER
ws["D3"] = "اكتب أي تاريخ بدل TODAY() لعرض يوم آخر"
ws["D3"].font = Font(name=FONT, italic=True, color="7F7F7F")


def day_cnt(status):
    return f'COUNTIFS({att("A")},$B$3,{att("P")},"{status}")'


def kpi(ref_lab, ref_val, text, formula, fmt="0"):
    c = ws[ref_lab]
    c.value, c.font, c.fill, c.border, c.alignment = text, F_HEAD, FILL_HEAD_CALC, BORDER, CENTER
    v = ws[ref_val]
    v.value, v.font, v.fill, v.border, v.alignment, v.number_format = (
        formula, Font(name=FONT, size=14, bold=True), FILL_KPI, BORDER, CENTER, fmt)


ws["A5"] = "مؤشرات اليوم"
ws["A5"].font = F_BOLD
kpi("A6", "B6", "الموظفين النشطين", f'=COUNTIF({q(S_EMP)}!$J$2:$J${EMP_LAST},"نشط")')
kpi("D6", "E6", "مسجلين في اليوم", f'=COUNTIF({att("A")},$B$3)')
kpi("G6", "H6", "حاضر", "=" + day_cnt("حاضر") + "+" + day_cnt("بدون انصراف"))
kpi("A7", "B7", "متأخر", "=" + day_cnt("متأخر"))
kpi("D7", "E7", "غائب", "=" + day_cnt("غائب") + "+" + day_cnt("غياب"))
kpi("G7", "H7", "إجازة / مأمورية",
    f'=COUNTIFS({att("A")},$B$3,{att("P")},"إجازة*")+COUNTIFS({att("A")},$B$3,{att("P")},"مأمورية*")'
    f'+COUNTIFS({att("A")},$B$3,{att("P")},"مهمة رسمية")+COUNTIFS({att("A")},$B$3,{att("P")},"غياب بعذر")')
kpi("A8", "B8", "بدون انصراف", "=" + day_cnt("بدون انصراف"))
kpi("D8", "E8", "نسبة الحضور", '=IF(E6=0,0,(H6+B7)/E6)', "0.0%")
kpi("G8", "H8", "لم يُسجَّل لهم اليوم", "=MAX(0,B6-E6)")

ws["A10"] = "مؤشرات الشهر"
ws["A10"].font = F_BOLD
ws["B10"] = f"={SET}!B3"
ws["B10"].number_format = FMT_MONTH
ws["B10"].font = F_LINK
P = q(S_PAY)
kpi("A11", "B11", "إجمالي أيام الحضور", f"={P}!D{PAY_TOTAL_ROW}")
kpi("D11", "E11", "إجمالي أيام الغياب", f"={P}!F{PAY_TOTAL_ROW}")
kpi("G11", "H11", "إجمالي أيام الإجازة", f"={P}!G{PAY_TOTAL_ROW}")
kpi("A12", "B12", "إجمالي ساعات العمل", f"={P}!H{PAY_TOTAL_ROW}", FMT_NUM)
kpi("D12", "E12", "إجمالي ساعات الإضافي", f"={P}!K{PAY_TOTAL_ROW}", FMT_NUM)
kpi("G12", "H12", "إجمالي دقائق التأخير", f"={P}!I{PAY_TOTAL_ROW}")
kpi("A13", "B13", "قيمة الإضافي", f"={P}!N{PAY_TOTAL_ROW}", FMT_MONEY)
kpi("D13", "E13", "إجمالي الخصومات", f"={P}!O{PAY_TOTAL_ROW}+{P}!P{PAY_TOTAL_ROW}", FMT_MONEY)
kpi("G13", "H13", "صافي الرواتب", f"={P}!Q{PAY_TOTAL_ROW}", FMT_MONEY)
for r in (6, 7, 8, 11, 12, 13):
    ws.row_dimensions[r].height = 30

# sheet order: help, dashboard, data entry, reports, setup
order = [S_HELP, S_DASH, S_ATT, S_EMP, S_LV, S_MS, S_PAY, S_REP, S_SHIFT, S_SET]
wb._sheets = [wb[n] for n in order]
tab = {S_HELP: "7F7F7F", S_DASH: "1F3864", S_ATT: "C55A11", S_EMP: "C55A11", S_LV: "C55A11",
       S_MS: "C55A11", S_PAY: "2E75B6", S_REP: "2E75B6", S_SHIFT: "548235", S_SET: "548235"}
for n, color in tab.items():
    wb[n].sheet_properties.tabColor = color
wb.active = 2
wb.save(OUT)
print("saved", OUT)
